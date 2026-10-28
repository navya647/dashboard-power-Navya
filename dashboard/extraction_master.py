"""DISCOM_SoP_master.xlsx -> ui/public/data/sop_master.json

One output row per sheet row, in sheet order. Every cell is written as Excel
displays it (number formats applied); text is copied unchanged. The only added
field is `standardised_num`, a plain number for plotting, taken from
`standardised_data` when that cell is numeric or a clean number/percentage.
"""
import json
import re
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent
SRC = ROOT.parent / "DISCOM_SoP_master.xlsx"
OUT = ROOT / "ui" / "public" / "data" / "sop_master.json"

COLUMNS = [
    "id", "state", "discom_code", "discom_name", "financial_year", "category",
    "indicator_family_id", "indicator_family", "unit", "indicator", "definition",
    "standard", "reported_data", "standardised_data", "comparison_possible",
    "standard_met", "standard_met_q1", "standard_met_q2", "standard_met_q3",
    "standard_met_q4", "standard_met_detail", "why_not_comparable",
    "interpretation", "documentation", "source_file", "applicable_regulation",
    "reported_data_source",
]

NUM_TEXT = re.compile(r"^-?\d+(?:\.\d+)?%?$")


def general(v):
    # Excel "General": integers without a decimal point, otherwise up to 10 significant digits
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, int) or float(v).is_integer():
        return str(int(v))
    return f"{v:.10g}"


def display(cell):
    v = cell.value
    if v is None:
        return None
    if isinstance(v, str):
        return v
    fmt = cell.number_format
    if fmt == "0.00%":
        return f"{v * 100:.2f}%"
    if fmt == "0%":
        return f"{v * 100:.0f}%"
    if fmt == "General":
        return general(v)
    raise ValueError(f"{cell.coordinate}: unhandled number format {fmt!r} for {v!r}")


def plot_number(cell):
    v = cell.value
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return round(v * 100, 10) if "%" in cell.number_format else v
    s = v.strip()
    if NUM_TEXT.match(s):
        return float(s.rstrip("%"))
    return None


def main():
    ws = openpyxl.load_workbook(SRC).active
    header = [c.value for c in ws[1]]
    if header != COLUMNS:
        raise SystemExit(f"Unexpected header: {header}")

    rows, state_order, states, families = [], [], {}, {}
    for cells in ws.iter_rows(min_row=2):
        if all(c.value is None for c in cells):
            continue
        r = {col: display(c) for col, c in zip(COLUMNS, cells)}
        r["standardised_num"] = plot_number(cells[COLUMNS.index("standardised_data")])
        rows.append(r)

        st = r["state"]
        if st not in states:
            state_order.append(st)
            states[st] = {"discoms": []}
        if not any(d["code"] == r["discom_code"] for d in states[st]["discoms"]):
            states[st]["discoms"].append({"code": r["discom_code"], "name": r["discom_name"]})
        # category is not constant within a family in the sheet (e.g. F014), so it stays per row only
        families.setdefault(r["indicator_family_id"], {
            "name": r["indicator_family"], "unit": r["unit"],
        })

    OUT.write_text(json.dumps({
        "source": SRC.name,
        "columns": COLUMNS,
        "state_order": state_order,
        "states": states,
        "families": families,
        "rows": rows,
    }, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print(f"{len(rows)} rows, {len(states)} states, {len(families)} families -> {OUT}")


if __name__ == "__main__":
    main()
